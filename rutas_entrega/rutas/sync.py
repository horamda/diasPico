"""Carga de datos: semilla, importación del Excel y sincronización con ChessERP."""
import json
import os
from datetime import datetime

from . import dias as D
from .db import Session
from .models import Cliente, ClienteEntrega, Deposito, Localidad, PlanRuta, SyncLog, Vehiculo

HERE = os.path.dirname(__file__)


# ---------- semilla (depósitos, localidades, vehículos y plan de la hoja INFORME) ----------
def cargar_semilla(path=None):
    s = Session()
    seed = json.load(open(path or os.path.join(HERE, "data", "seed.json"), encoding="utf-8"))
    deps = {}
    for d in seed["depositos"]:
        dep = s.query(Deposito).filter_by(nombre=d["nombre"]).one_or_none() or Deposito(nombre=d["nombre"])
        dep.direccion, dep.lat, dep.lng = d.get("direccion"), d.get("lat"), d.get("lng")
        s.add(dep)
        deps[d["nombre"]] = dep
    s.flush()
    locs = {}
    for l in seed["localidades"]:
        loc = s.query(Localidad).filter_by(nombre=l["nombre"]).one_or_none() or Localidad(nombre=l["nombre"])
        loc.deposito_id = deps[l["sucursal"]].id
        loc.lat, loc.lng = loc.lat or l.get("lat"), loc.lng or l.get("lng")
        s.add(loc)
        locs[l["nombre"]] = loc
    s.flush()
    for v in seed["vehiculos"]:
        dep = deps[v["sucursal"]]
        veh = s.query(Vehiculo).filter_by(deposito_id=dep.id, nombre=v["nombre"]).one_or_none()
        if veh:
            continue  # no pisar un plan ya editado
        veh = Vehiculo(deposito_id=dep.id, nombre=v["nombre"], color=v.get("color", "#c4561e"))
        s.add(veh)
        s.flush()
        for dia, lista in v["plan"].items():
            for i, nombre in enumerate(lista, 1):
                s.add(PlanRuta(vehiculo_id=veh.id, dia=dia, localidad_id=locs[nombre].id, orden=i))
    s.commit()


def _localidad_por_nombre(s, nombre, cache):
    if not nombre:
        return None
    if nombre not in cache:
        cache[nombre] = s.query(Localidad).filter_by(nombre=nombre).one_or_none()
    return cache[nombre]


# ---------- importación del Excel "ARMADO DE DIAS DE ENTREGA" ----------
def importar_excel(path, sucursal_erp=2, usuario="import-excel", pisar_dias=False):
    """Lee las hojas CLIENTES, LOCALIDAD-SUCURSAL y RUTAS DE ENTREGA.

    - CLIENTES (filtrado por sucursal_erp) -> rt_cliente
    - LOCALIDAD-SUCURSAL -> rt_localidad (crea las que falten)
    - RUTAS DE ENTREGA -> rt_cliente_entrega.dias (no pisa días ya editados salvo pisar_dias=True)
    Nota: en RUTAS DE ENTREGA las columnas LAT/LONG vienen invertidas; acá se usan las del ERP.
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    s = Session()
    log = SyncLog(origen="excel")
    s.add(log)

    def filas(nombre):
        ws = wb[nombre]
        it = ws.iter_rows(values_only=True)
        head = [str(h).strip() if h is not None else "" for h in next(it)]
        for row in it:
            yield dict(zip(head, row))

    # localidades -> depósito
    if "LOCALIDAD-SUCURSAL" in wb.sheetnames:
        deps = {d.nombre: d for d in s.query(Deposito)}
        for r in filas("LOCALIDAD-SUCURSAL"):
            nom, suc = (r.get("LOCALIDAD") or "").strip().upper(), (r.get("SUCURSAL") or "").strip().upper()
            if not nom:
                continue
            dep = deps.get(suc) or Deposito(nombre=suc)
            if suc not in deps:
                s.add(dep); s.flush(); deps[suc] = dep
            loc = s.query(Localidad).filter_by(nombre=nom).one_or_none() or Localidad(nombre=nom)
            loc.deposito_id = dep.id
            s.add(loc)
        s.flush()

    ahora = datetime.utcnow()
    existentes = {c.id_cliente: c for c in s.query(Cliente)}
    for r in filas("CLIENTES"):
        if r.get("Sucursal") != sucursal_erp or r.get("Cliente") is None:
            continue
        cid = int(r["Cliente"])
        lat = r.get("Coord Y de entrega") or r.get("Coord Y")
        lng = r.get("Coord X de entrega") or r.get("Coord X")
        datos = dict(
            id_sucursal_erp=sucursal_erp, razon_social=r.get("Razon social"), fantasia=r.get("Nombre de fantasia"),
            domicilio=r.get("Domicilio"), localidad_erp=(r.get("Descripcion localidad") or "").strip().upper() or None,
            lat=float(lat) if lat else None, lng=float(lng) if lng else None,
            dias_erp=D.desde_texto_erp(r.get("Fuerza de venta 1 Dias de entrega")),
            vendedor=r.get("Fuerza de venta 1 Descripcion personal comercial"),
            ruta_venta=str(r.get("Fuerza de venta 1 Ruta de venta") or "").replace(".0", "") or None,
            ruta_distribucion=r.get("Fuerza de venta 1 Descripcion ruta de distribucion"),
            horario=r.get("Horario de entrega"), canal=r.get("Descripcion subcanal"),
            anulado=(str(r.get("Anulado") or "").upper() == "SI"), sincronizado=ahora)
        c = existentes.get(cid)
        if c is None:
            c = Cliente(id_cliente=cid); s.add(c); existentes[cid] = c; log.nuevos += 1
        else:
            log.actualizados += 1
        for k, v in datos.items():
            setattr(c, k, v)
        log.leidos += 1
    s.flush()

    # días de la planilla
    cols = {"LU": "LUNES", "MA": "MARTES", "MI": "MIERCOLES", "JU": "JUEVES", "VI": "VIERNES", "SA": "SABADO"}
    en_planilla = set()
    for r in filas("RUTAS DE ENTREGA"):
        cid = r.get("CLIENTE")
        if cid is None or int(cid) not in existentes:
            continue
        cid = int(cid)
        en_planilla.add(cid)
        dias = "".join(d for d, col in cols.items() if r.get(col) or r.get(col + " "))
        e = s.get(ClienteEntrega, cid)
        if e is None:
            s.add(ClienteEntrega(id_cliente=cid, dias=dias, en_planilla=True, actualizado=ahora, actualizado_por=usuario))
        else:
            e.en_planilla = True
            if pisar_dias:
                e.dias = dias
    for cid in existentes:
        if cid not in en_planilla and s.get(ClienteEntrega, cid) is None:
            s.add(ClienteEntrega(id_cliente=cid, dias="", en_planilla=False, actualizado=ahora, actualizado_por=usuario))
    log.fin = datetime.utcnow()
    s.commit()
    return {"leidos": log.leidos, "nuevos": log.nuevos, "actualizados": log.actualizados, "en_planilla": len(en_planilla)}


# ---------- sincronización con ChessERP ----------
def sincronizar_chess(sucursales=None, client=None, inicializar_dias_desde_erp=False):
    """Trae los clientes de Chess y actualiza rt_cliente. No toca rt_cliente_entrega,
    salvo para crear la fila de clientes nuevos (con los días del ERP si se pide)."""
    from .chess import ChessClient, mapear_cliente
    sucursales = sucursales or [int(x) for x in os.environ.get("CHESS_SUCURSALES", "2").split(",")]
    client = client or ChessClient()
    s = Session()
    log = SyncLog(origen="chess")
    s.add(log); s.commit()
    try:
        client.login()
        ahora = datetime.utcnow()
        for suc in sucursales:
            for raw in client.todos_los_clientes(suc):
                d = mapear_cliente(raw)
                c = s.get(Cliente, d["id_cliente"])
                if c is None:
                    c = Cliente(id_cliente=d["id_cliente"]); s.add(c); log.nuevos += 1
                    s.add(ClienteEntrega(id_cliente=d["id_cliente"],
                                         dias=d["dias_erp"] if inicializar_dias_desde_erp else "",
                                         en_planilla=False, actualizado=ahora, actualizado_por="sync-chess"))
                else:
                    log.actualizados += 1
                for k, v in d.items():
                    setattr(c, k, v)
                c.sincronizado = ahora
                log.leidos += 1
                if log.leidos % 500 == 0:
                    s.flush()
        log.fin = datetime.utcnow()
        s.commit()
    except Exception as e:  # se registra y se relanza
        s.rollback()
        log = s.merge(log)
        log.error, log.fin = repr(e)[:2000], datetime.utcnow()
        s.commit()
        raise
    return {"leidos": log.leidos, "nuevos": log.nuevos, "actualizados": log.actualizados}
